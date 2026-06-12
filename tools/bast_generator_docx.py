#!/usr/bin/env python3
"""Generate a BAST DOCX by replacing {{PLACEHOLDER}} markers."""

import argparse
import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = ROOT / "templates/bast_vapt_template.docx"
DEFAULT_DATA = ROOT / "generated_reports/bast_data.json"
DEFAULT_OUTPUT = ROOT / "generated_reports/bast.docx"


def marker_variants(key: str) -> set[str]:
    normalized = key.strip().upper()
    variants = {normalized}
    variants.add(normalized.replace(" ", "_"))
    variants.add(normalized.replace("_", " "))
    return {f"{{{{{variant}}}}}" for variant in variants if variant}


def build_replacements(data: dict[str, Any]) -> dict[str, str]:
    replacements: dict[str, str] = {}
    for key, value in data.items():
      if value is None:
          value = ""
      text = str(value)
      for marker in marker_variants(key):
          replacements[marker] = text
    return replacements


def replace_in_xml_element(element: Any, replacements: dict[str, str]) -> None:
    for paragraph in element.iter(qn("w:p")):
        text_nodes = list(paragraph.iter(qn("w:t")))
        if not text_nodes:
            continue
        for marker, value in replacements.items():
            replace_marker_in_text_nodes(text_nodes, marker, value)


def replace_marker_in_text_nodes(text_nodes: list[Any], marker: str, value: str) -> None:
    # Fast path: most template placeholders live inside a single Word text node.
    for node in text_nodes:
        if node.text and marker in node.text:
            node.text = node.text.replace(marker, value)

    # Fallback for placeholders split across Word runs. This removes only the
    # marker characters and keeps surrounding runs/tabs/formatting intact.
    combined = "".join(node.text or "" for node in text_nodes)
    starts: list[int] = []
    pos = combined.find(marker)
    while pos >= 0:
        starts.append(pos)
        pos = combined.find(marker, pos + len(marker))

    for start in reversed(starts):
        end = start + len(marker)
        spans = []
        cursor = 0
        for node in text_nodes:
            text = node.text or ""
            node_start = cursor
            node_end = cursor + len(text)
            if node_end > start and node_start < end:
                spans.append((node, text, node_start, node_end))
            cursor = node_end

        if not spans:
            continue

        first_node, first_text, first_start, first_end = spans[0]
        prefix = first_text[: max(0, start - first_start)]
        suffix = first_text[max(0, end - first_start) :] if end <= first_end else ""
        first_node.text = prefix + value + suffix

        for node, text, node_start, node_end in spans[1:]:
            keep_prefix = text[: max(0, start - node_start)] if start > node_start else ""
            keep_suffix = text[max(0, end - node_start) :] if end < node_end else ""
            node.text = keep_prefix + keep_suffix


def unresolved_markers(doc: Document) -> list[str]:
    texts: list[str] = []
    elements = [doc.element.body]
    for section in doc.sections:
        elements.append(section.header._element)
        elements.append(section.footer._element)
    for element in elements:
        texts.extend(node.text or "" for node in element.iter(qn("w:t")))
    return sorted(set(re.findall(r"\{\{[^{}]+\}\}", "\n".join(texts))))


def generate(template: Path, data_path: Path, output: Path) -> dict[str, Any]:
    data = json.loads(data_path.read_text())
    doc = Document(template)
    replacements = build_replacements(data)

    replace_in_xml_element(doc.element.body, replacements)
    for section in doc.sections:
        replace_in_xml_element(section.header._element, replacements)
        replace_in_xml_element(section.footer._element, replacements)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    unresolved = unresolved_markers(doc)
    return {
        "output": str(output),
        "unresolved": unresolved,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate a BAST DOCX from placeholder data.")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--data", type=Path, default=DEFAULT_DATA)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = generate(args.template, args.data, args.out)
    print(json.dumps(result, indent=2))
    return 0 if not result["unresolved"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
